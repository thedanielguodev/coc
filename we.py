#!/usr/bin/env python3
"""
Download California Continuum of Care (CoC) data from HUD:
  - CoC administrative boundaries (GeoJSON, WGS84)
  - Point-in-Time (PIT) homeless counts by CoC, 2007-2024
  - Housing Inventory Count (HIC) bed/unit inventory by CoC, 2007-2024
  - PIT veteran counts by CoC, 2011-2024

Sources (all official HUD):
  Boundaries : HUD-eGIS Open Data (hudgis-hud.opendata.arcgis.com)
               "Continuum of Care (CoC) Grantee Areas"
  PIT/HIC    : HUD USER AHAR data downloads (huduser.gov)

Output is written to ./data/ as GeoJSON + CSV.
"""

import csv
import json
import math
import pathlib

import requests

OUTPUT_DIR = pathlib.Path(__file__).parent / "data"

UA = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
)

BOUNDARIES_URL = (
    "https://hudgis-hud.opendata.arcgis.com/api/download/v1/items/"
    "c930d736b1764c259371fc7111e02740/geojson?layers=0"
)

# huduser.gov blocks requests without a Referer from the AHAR page.
HUDUSER_REFERER = (
    "https://www.huduser.gov/portal/datasets/ahar/"
    "2024-ahar-part-1-pit-estimates-of-homelessness-in-the-us.html"
)
HUDUSER_BASE = "https://www.huduser.gov/portal/sites/default/files/xls/"

PIT_FILE = "2007-2024-PIT-Counts-by-CoC.xlsb"
HIC_FILE = "2007-2024-HIC-Counts-by-CoC.xlsx"
PIT_VET_FILE = "2011-2024-PIT-Veteran-Counts-by-CoC.xlsx"

# Headline PIT metrics that exist (by name) across most/all survey years.
# Anything missing in a given year's sheet is simply left blank.
CORE_PIT_METRICS = [
    "Overall Homeless",
    "Sheltered ES Homeless",
    "Sheltered TH Homeless",
    "Sheltered SH Homeless",
    "Sheltered Total Homeless",
    "Unsheltered Homeless",
    "Overall Homeless Individuals",
    "Sheltered Total Homeless Individuals",
    "Unsheltered Homeless Individuals",
    "Overall Homeless People in Families",
    "Sheltered Total Homeless People in Families",
    "Unsheltered Homeless People in Families",
    "Overall Homeless Family Households",
    "Sheltered Total Homeless Family Households",
    "Unsheltered Homeless Family Households",
    "Overall Chronically Homeless Individuals",
    "Sheltered Total Chronically Homeless Individuals",
    "Unsheltered Chronically Homeless Individuals",
    "Overall Homeless Veterans",
    "Sheltered Total Homeless Veterans",
    "Unsheltered Homeless Veterans",
    "Overall Homeless Unaccompanied Youth (Under 25)",
    "Sheltered Total Homeless Unaccompanied Youth (Under 25)",
    "Unsheltered Homeless Unaccompanied Youth (Under 25)",
]

CORE_HIC_METRICS = [
    "Total Year-Round Beds (ES, TH, SH)",
    "Total HMIS Year-Round Beds (ES, TH, SH)",
    "Total Year-Round Beds (ES)",
    "Total Year-Round Beds (TH)",
    "Total Year-Round Beds (SH)",
    "Total Year-Round Beds (RRH)",
    "Total Year-Round Beds (PSH)",
    "Total Year-Round Beds (OPH)",
    "Dedicated Veteran Beds (ES, TH, SH)",
    "Dedicated Youth Beds (ES, TH, SH)",
]


def web_mercator_to_wgs84(coords):
    """Recursively reproject nested GeoJSON coordinate arrays EPSG:3857 -> EPSG:4326."""
    if isinstance(coords[0], (int, float)):
        x, y = coords[0], coords[1]
        lon = x / 20037508.34 * 180
        lat = 180 / math.pi * (2 * math.atan(math.exp(y / 20037508.34 * math.pi)) - math.pi / 2)
        return [lon, lat]
    return [web_mercator_to_wgs84(c) for c in coords]


def fetch_boundaries():
    print("Fetching CoC boundaries from HUD-eGIS Open Data...")
    resp = requests.get(BOUNDARIES_URL, headers={"User-Agent": UA}, timeout=60)
    resp.raise_for_status()
    data = resp.json()

    ca_features = [f for f in data["features"] if str(f["properties"].get("COCNUM", "")).startswith("CA-")]
    for feature in ca_features:
        feature["geometry"]["coordinates"] = web_mercator_to_wgs84(feature["geometry"]["coordinates"])

    out = {
        "type": "FeatureCollection",
        "crs": {"type": "name", "properties": {"name": "EPSG:4326"}},
        "features": ca_features,
    }

    dest = OUTPUT_DIR / "ca_coc_boundaries.geojson"
    dest.write_text(json.dumps(out))
    print(f"  {len(ca_features)} CA CoCs -> {dest}")
    return {f["properties"]["COCNUM"] for f in ca_features}


def download_huduser_file(filename):
    dest = OUTPUT_DIR / "_raw" / filename
    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.exists():
        return dest
    print(f"Downloading {filename}...")
    resp = requests.get(
        HUDUSER_BASE + filename,
        headers={"User-Agent": UA, "Referer": HUDUSER_REFERER},
        timeout=120,
    )
    resp.raise_for_status()
    dest.write_bytes(resp.content)
    return dest


def parse_pit(path, ca_codes):
    import pyxlsb

    rows = []
    with pyxlsb.open_workbook(str(path)) as wb:
        for sheet_name in wb.sheets:
            if not sheet_name.isdigit():
                continue
            year = int(sheet_name)
            with wb.get_sheet(sheet_name) as sheet:
                it = sheet.rows()
                header = [c.v for c in next(it)]
                col = {name: i for i, name in enumerate(header) if name is not None}
                if "CoC Number" not in col:
                    continue
                for r in it:
                    vals = [c.v for c in r]
                    coc_number = vals[col["CoC Number"]]
                    if coc_number not in ca_codes:
                        continue
                    name_idx = col.get("CoC Name")
                    record = {
                        "coc_number": coc_number,
                        "coc_name": vals[name_idx] if name_idx is not None else None,
                        "year": year,
                    }
                    for metric in CORE_PIT_METRICS:
                        idx = col.get(metric)
                        record[metric] = vals[idx] if idx is not None else None
                    rows.append(record)
    return rows


def _norm_header(name):
    """Collapse whitespace so equivalent headers match across years, e.g.
    'Total Year-Round Beds (ES, TH, SH)' (2014+) vs '...(ES,TH,SH)' (2008-2012).
    Does NOT collapse genuinely different column sets (e.g. 2013's
    '(ES,TH,RRH,SH)' still won't match '(ES,TH,SH)') so it never silently
    merges two differently-defined metrics.
    """
    return "".join(str(name).split()) if name is not None else name


def parse_hic(path, ca_codes):
    import openpyxl

    rows = []
    wb = openpyxl.load_workbook(str(path), read_only=True)
    for sheet_name in wb.sheetnames:
        if not sheet_name.isdigit():
            continue
        year = int(sheet_name)
        ws = wb[sheet_name]
        it = ws.iter_rows(values_only=True)
        next(it)  # group-header row
        header = list(next(it))
        col = {}
        for i, name in enumerate(header):
            key = _norm_header(name)
            if key is not None and key not in col:
                col[key] = i
        # 2007-2012 HIC sheets label the CoC ID column "CoC" instead of "CoC Number".
        coc_idx = col.get(_norm_header("CoC Number"))
        if coc_idx is None:
            coc_idx = col.get(_norm_header("CoC"))
        if coc_idx is None:
            continue
        for r in it:
            coc_number = r[coc_idx]
            if coc_number not in ca_codes:
                continue
            record = {
                "coc_number": coc_number,
                "year": year,
            }
            for metric in CORE_HIC_METRICS:
                idx = col.get(_norm_header(metric))
                record[metric] = r[idx] if idx is not None else None
            rows.append(record)
    return rows


def parse_pit_veterans(path, ca_codes):
    import openpyxl
    import re

    rows = []
    wb = openpyxl.load_workbook(str(path), read_only=True)
    for sheet_name in wb.sheetnames:
        if not sheet_name.isdigit():
            continue
        year = int(sheet_name)
        ws = wb[sheet_name]
        it = ws.iter_rows(values_only=True)
        header = [re.sub(r",\s*\d{4}$", "", str(h)) if h else h for h in next(it)]
        col = {name: i for i, name in enumerate(header) if name is not None}
        if "CoC Number" not in col:
            continue
        for r in it:
            coc_number = r[col["CoC Number"]]
            if coc_number not in ca_codes:
                continue
            record = {"coc_number": coc_number, "year": year}
            for name, idx in col.items():
                if name in ("CoC Number",):
                    continue
                record[name] = r[idx]
            rows.append(record)
    return rows


def write_csv(rows, fieldnames, dest):
    with open(dest, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
    print(f"  {len(rows)} rows -> {dest}")


def main():
    OUTPUT_DIR.mkdir(exist_ok=True)

    ca_codes = fetch_boundaries()

    pit_path = download_huduser_file(PIT_FILE)
    hic_path = download_huduser_file(HIC_FILE)
    vet_path = download_huduser_file(PIT_VET_FILE)

    print("Parsing PIT counts (2007-2024)...")
    pit_rows = parse_pit(pit_path, ca_codes)
    write_csv(
        pit_rows,
        ["coc_number", "coc_name", "year"] + CORE_PIT_METRICS,
        OUTPUT_DIR / "ca_coc_pit_by_year.csv",
    )

    print("Parsing HIC bed inventory (2007-2024)...")
    hic_rows = parse_hic(hic_path, ca_codes)
    write_csv(
        hic_rows,
        ["coc_number", "year"] + CORE_HIC_METRICS,
        OUTPUT_DIR / "ca_coc_hic_by_year.csv",
    )

    print("Parsing PIT veteran counts (2011-2024)...")
    vet_rows = parse_pit_veterans(vet_path, ca_codes)
    vet_fields = ["coc_number", "year"] + [k for k in vet_rows[0] if k not in ("coc_number", "year")]
    write_csv(vet_rows, vet_fields, OUTPUT_DIR / "ca_coc_pit_veterans_by_year.csv")

    print(f"\nDone. All CA CoC data written to {OUTPUT_DIR}/")


if __name__ == "__main__":
    main()
